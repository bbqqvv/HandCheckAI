from datetime import datetime
import openpyxl
from openpyxl.styles import Font, Alignment


class MeasurementExporter:
    """Handles exporting measurement data to Excel."""

    def __init__(self):
        self.data = []

    def add_measurement(self, image_name, hand_id, angle):
        """Add a new measurement to the collection."""
        self.data.append({
            'image_name': image_name,
            'hand_id': hand_id,
            'angle': angle,
            'timestamp': datetime.now()
        })

    def get_interpretation(self, angle):
        """Get interpretation text for the angle."""
        if abs(angle) < 15:
            return "Nearly vertical"
        elif angle > 0:
            return f"Tilted {angle:.1f}° to the right"
        else:
            return f"Tilted {abs(angle):.1f}° to the left"

    def export_to_excel(self, file_path):
        """Export all measurements to an Excel file."""
        if not self.data:
            return False

        try:
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Hand Measurements"

            # Write headers
            headers = ["Image Name", "Hand", "Angle (Degrees)", "Interpretation", "Timestamp"]
            ws.append(headers)

            # Format headers
            for cell in ws[1]:
                cell.font = Font(bold=True)
                cell.alignment = Alignment(horizontal='center')

            # Write data
            for measurement in self.data:
                ws.append([
                    measurement['image_name'],
                    measurement['hand_id'],
                    measurement['angle'],
                    self.get_interpretation(measurement['angle']),
                    measurement['timestamp'].strftime('%Y-%m-%d %H:%M:%S')
                ])

            # Format columns
            for col in ws.columns:
                max_length = 0
                column = col[0].column_letter

                for cell in col:
                    try:
                        if len(str(cell.value)) > max_length:
                            max_length = len(str(cell.value))
                    except:
                        pass

                adjusted_width = (max_length + 2) * 1.2
                ws.column_dimensions[column].width = adjusted_width

            # Format angle column
            for row in ws.iter_rows(min_row=2, max_col=3, max_row=ws.max_row):
                for cell in row[2:3]:
                    cell.number_format = '0.00'

            # Save file
            wb.save(file_path)
            return True

        except Exception as e:
            print(f"Error saving Excel file: {e}")
            return False